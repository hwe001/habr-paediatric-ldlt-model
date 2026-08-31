import docx

d = docx.Document()
d.add_heading("Response to Reviewer, Round 4", level=1)
d.add_paragraph(
    "We thank the reviewer for identifying the equation notation issue "
    "and the incomplete description of the arterial-growth analysis; "
    "both are addressed below, along with the wording and reporting "
    "issues raised."
)

points = [
    ("1. Factor-of-60 equation error",
     "Verified directly against the implementation: the code was already "
     "correct (mean HA flow 25.0 mL/min = 0.417 mL/s; dividing by "
     "A_HA=0.0113 cm^2 gives 36.9 cm/s, matching your own back-of-"
     "envelope check and consistent with the simulated PSV of 53.1 "
     "cm/s). The error was in the manuscript's written equation only. "
     "Corrected to v_HA(tau) = Q_HA(tau)/A_HA (Section 2.2.1), with "
     "Section 2.2 now stating explicitly that Q is in mL/s during "
     "integration, and a one-line dimensional check added."),
    ("2. HA growth must state what changes in the circuit",
     "Agreed, and implemented as you suggested: the primary analysis is "
     "now explicitly labelled KINEMATIC DILUTION ONLY (only A_HA, the "
     "flow-to-velocity conversion area, changes; Rs_HA is still solved "
     "via the same HABR flow-target procedure, L_HA is unchanged) -- "
     "Section 2.7 and 3.5 state this precisely, and no longer use the "
     "phrase 'calibre-growth sensitivity dominates' without this "
     "qualification. A second, EXPLORATORY STRUCTURAL REMODELLING "
     "variant was added (new Table 4): L_HA additionally scaled "
     "geometrically (~1/(1+g)^2) with assumed growth, Rs_HA still solved "
     "via the HABR flow target using this grown L_HA. The two give "
     "quantitatively close results (e.g. 191.2% vs. 185.6% at 40% "
     "growth) -- reported as a genuine check, not assumed to agree in "
     "advance."),
    ("3. 'Entire admissible parameter family' overstated",
     "Corrected throughout to name the explored range explicitly "
     "(P_HA_amp in [23.5, 60] mmHg). We tested extending to 300 mmHg "
     "(a supraphysiological pulse pressure, for mathematical convergence "
     "only) and found the envelope stabilises (20.2% at 60 mmHg to 18.8% "
     "at 300 mmHg) -- both the finding and this practical justification "
     "for the explored range are now stated (Section 2.7, 3.3). The full "
     "grid (19 points), L_HA range (2.16-22.23), Rs_HA range (127.92-"
     "127.93), fitting tolerance (objective <1e-6), and root-finding "
     "details (bounds, expansion procedure, failure criterion) are now "
     "reported (Section 2.3, 2.7, Table 3)."),
    ("4. Calibre-growth ranges called 'plausible' without support",
     "Changed to 'illustrative' throughout (Sections 2.7, 3.5, Abstract, "
     "Discussion, Conclusion), with your suggested phrasing adopted "
     "closely: we state explicitly that no source measures serial "
     "calibre change in this population, and that this is precisely the "
     "gap the analysis is designed to expose."),
    ("5. Baseline HA diameter combined with growth",
     "Added (Section 3.5): growth sensitivity (0-40%) repeated at "
     "baseline diameters 1.0, 1.2, and 1.4 mm. The growth-sensitivity "
     "shape is consistent across all three (e.g. at 20% growth: 129.5%, "
     "128.8%, 127.7% respectively) -- the conclusion does not depend on "
     "which value within the published uncertainty is assumed."),
]
for title, body in points:
    d.add_heading(title, level=3)
    d.add_paragraph(body)

d.add_heading("Interpretation and presentation", level=3)
d.add_paragraph(
    "Abstract reworded to state precisely which parameter is stable "
    "(Rs_HA) versus which two trade off (L_HA, P_HA_amp). Section 3.3 "
    "now explains why the fraction reproduced still varies (20.2-27.1%) "
    "despite near-constant Rs_HA and mean flow: PSV depends on the "
    "pulsatile shape of the waveform, not only its mean, so different "
    "(L_HA, P_HA_amp) combinations respond somewhat differently once "
    "Rs_HA is re-solved to a new target at POD30. 'Explained' changed to "
    "'reproduced' throughout Section 3.5. Section 3.8's phrasing "
    "(allowing model misspecification and measurement error) is "
    "unchanged, as you noted it was already correct."
)
d.add_paragraph(
    "On the code repository: we have not yet created one, and would "
    "rather confirm this with the corresponding author before doing so "
    "than state a URL that does not yet exist. The Data/code section "
    "now flags this explicitly as a pre-submission action item rather "
    "than implying it is already arranged."
)

d.save("Response_to_Reviewer_Round4.docx")
print("Saved: Response_to_Reviewer_Round4.docx")
