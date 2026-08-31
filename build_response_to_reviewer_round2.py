import docx

d = docx.Document()
d.add_heading("Response to Reviewer, Round 2", level=1)
d.add_paragraph(
    "We thank the reviewer for identifying the dimensional inconsistency "
    "and the statistical error below; both required substantive correction, "
    "not only rewording."
)

points = [
    ("1. Dimensional inconsistency (Q_HA vs. Q_PV)",
     "Corrected, not merely disclosed. Q_HA is now a genuine mL/min-scale "
     "volumetric flow throughout the model, on the same physical basis as "
     "Q_PV, related to Doppler velocity only via an explicit hepatic-"
     "artery cross-sectional-area assumption (new Section 2.2.1): "
     "diameter 1.2 mm, from a healthy infant control group in a "
     "paediatric ultrasound reference study (Kim et al., 2007) -- the "
     "best published paediatric value we located, explicitly flagged as "
     "an assumption for this cohort's graft artery, not a measurement of "
     "it. All post-LDLT calculations were rerun under this corrected "
     "formulation. The diameter-sensitivity finding (fragility to a 10% "
     "portal-diameter-growth assumption) persists after the fix, "
     "confirming it is a real property of the model's response, not an "
     "artefact of the prior inconsistency. We no longer describe HABR as "
     "'directionally necessary'; Results 3.5 and the Conclusion now use "
     "language close to your suggested phrasing directly."),
    ("2. RI confidence-interval error",
     "You are correct and we have corrected the error throughout "
     "(Results 3.4, and the response text below): a 95% CI of "
     "[-0.010, +0.010] EXCLUDES a true difference of 0.020; it does not "
     "fail to exclude it. The manuscript now states the narrower, "
     "correct conclusion: no statistically detectable difference; "
     "compatible with true differences up to about 0.01; the clinical/"
     "mechanistic significance of a 0.01 change is a separate, "
     "unresolved question; and explicitly, this does NOT mean a "
     "0.02-magnitude change remains possible at POD7-30."),
    ("3. 'Equivalence test' mislabelled",
     "Removed throughout. Section 2.6 now describes this as an ordinary "
     "unpaired comparison of summary statistics (Welch's t-test and CI), "
     "with no equivalence margin or TOST procedure claimed."),
    ("4. Portal pathway / time-notation ambiguity",
     "Clarified explicitly in revised Methods 2.2/2.2.1/2.4: two distinct "
     "time variables are now named separately (tau for time within a "
     "cardiac cycle; 'day'/POD for postoperative time indexing "
     "independent simulations). The text states plainly that the "
     "portal-vein ODE is disabled in the post-LDLT configuration and "
     "Q_PV is a prescribed, day-specific constant, overwritten between "
     "days, not evolved continuously."),
    ("5. HABR-to-resistance update underdefined",
     "Added the exact update procedure (Methods 2.3): reference day, sign "
     "convention, the POD1-referenced (non-cumulative) convention used in "
     "the continuous scenario, the single-application convention in the "
     "acute-then-frozen scenario, and the root-finding procedure "
     "(Brent's method) with which parameters are held fixed during it."),
    ("6. Monte Carlo interval reconsideration",
     "Addressed directly (Methods 2.8, renamed 'illustrative sensitivity "
     "intervals'): states explicitly that POD1 and later means were "
     "sampled independently (the wrong covariance structure for this "
     "repeated-measures cohort), that no draws were discarded (including "
     "near-zero or negative measured-decline draws, which produce the "
     "heavy-tailed POD7 interval), that PVV and model-parameter "
     "uncertainty were not propagated, and that these are not proper "
     "95% confidence intervals."),
    ("7. Conclusions stronger than justified",
     "Revised throughout, largely adopting your suggested wording: "
     "Results 3.5, Discussion, and Conclusion now state that under the "
     "primary scenario's specific assumptions the model predicts a "
     "reduction, and that this direction and magnitude are not robust to "
     "plausible alternative portal-calibre assumptions -- 'directionally "
     "necessary' and similar phrases have been removed."),
]
for title, body in points:
    d.add_heading(title, level=3)
    d.add_paragraph(body)

d.add_heading("8. Response-letter tone", level=3)
d.add_paragraph(
    "Revised accordingly in this letter."
)

d.save("Response_to_Reviewer_Round2.docx")
print("Saved: Response_to_Reviewer_Round2.docx")
