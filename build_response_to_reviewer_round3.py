import docx

d = docx.Document()
d.add_heading("Response to Reviewer, Round 3", level=1)
d.add_paragraph(
    "We thank the reviewer for the continued detailed engagement. This "
    "revision focuses on propagating the two structural uncertainties "
    "identified, as recommended, rather than adding further caveats."
)

points = [
    ("1. POD1 calibration non-identifiability",
     "We generated the admissible (Rs_HA, L_HA, P_HA_amp) family directly "
     "(scanning P_HA_amp from its minimum feasible value, ~23.5 mmHg, to "
     "60 mmHg, re-solving Rs_HA and L_HA for each) and propagated every "
     "member through POD30 (new Section 2.7 / 3.3, Figure 6). The "
     "resulting envelope is 20.2-27.1% of the measured decline "
     "reproduced -- narrower than might be feared, because Rs_HA and the "
     "resulting mean HA flow are nearly invariant across the family (L_HA "
     "and P_HA_amp trade off to preserve the same pulsatility ratio, but "
     "the mean-flow quantity driving the HABR update is well-constrained "
     "by the two POD1 targets). We report this as a genuine finding, not "
     "an assumption: this particular identifiability gap turns out not "
     "to be the dominant source of uncertainty, in contrast to the "
     "calibre assumptions below."),
    ("2. HA diameter needs its own sensitivity analysis",
     "Added (new Section 2.7/3.5, Figures 7-8). The HA diameter's own "
     "published uncertainty (1.0-1.4 mm), properly re-anchored at POD1 "
     "for each value, has little effect (23.7-26.4%) -- an earlier, "
     "buggy version of this specific check failed to re-anchor and gave "
     "a spurious large effect, which we caught and fixed before "
     "reporting. Postoperative HA diameter GROWTH, however, has a large "
     "effect in the OPPOSITE direction to portal growth (25.3% at 0% "
     "growth to 191.2% at 40% growth). A 2D grid varying both jointly "
     "(Figure 8) shows the combined range is -44% to +129% across a "
     "modest, jointly plausible 0-20% growth in each vessel -- we no "
     "longer imply calibre uncertainty is a portal-only concern."),
    ("3. Physical units incompletely specified",
     "Table 3 and Methods 2.2 now state explicit units for every element "
     "(P in mmHg, Q in mL/s during integration, tau in seconds, giving "
     "Rs/Rp in mmHg*s/mL, L in mmHg*s^2/mL, C in mL/mmHg), derived "
     "directly from the units already implicit in the code rather than "
     "newly invented."),
    ("4. Mean flow vs. Doppler PSV",
     "Addressed directly in new Section 2.2.1: v=Q/A gives a cross-"
     "sectional mean velocity, not the spatial-peak quantity PSV "
     "conventionally represents, and no HA-side profile factor "
     "analogous to Chen et al.'s portal 0.57 is applied. Because this is "
     "absorbed into the POD1 fit, we no longer describe the resulting "
     "25.0 mL/min mean flow as an external plausibility check; Results "
     "3.2 states this explicitly, and only relative (percent) flow "
     "changes and RI (both invariant to a fixed unmodelled profile "
     "factor) are used in the conclusions."),
    ("5. Portal-flow formula / graft-mass scaling",
     "Ran the requested sensitivity check (new Section 2.7/3.5): scaling "
     "Q_PV by 0.5-2.0x (as if graft mass were 200-50 g) changes the "
     "POD30 result only from 25.5% to 25.1%. We report this as a genuine, "
     "reassuring finding -- this specific ambiguity, while conceptually "
     "valid, is not a material source of uncertainty in this model, "
     "unlike the calibre-growth assumptions."),
    ("6. Revision-history narrative in the manuscript",
     "Removed throughout the manuscript (Abstract, Introduction, Section "
     "2.2.1, figure captions, and one further passage in Section 2.7 "
     "describing how a bug in an intermediate check was found and fixed "
     "-- all now state the corrected method directly, with the process "
     "narrative confined to these response letters). 'Independent "
     "literature targets' changed to 'previously-derived calibration "
     "targets' throughout."),
]
for title, body in points:
    d.add_heading(title, level=3)
    d.add_paragraph(body)

d.add_heading("Minor points", level=3)
d.add_paragraph(
    "Root-finding is now named precisely (Brent's method via "
    "scipy.optimize.brentq), with search bounds, bracket-expansion "
    "procedure, tolerance, and failure criterion stated (Section 2.3). "
    "The representative POD1 solution's origin is now explained "
    "(Section 2.4: the initial guess and warm-start procedure used to "
    "generate the admissible family, of which the representative member "
    "is simply the first solved). 'Whatever the model does not reproduce "
    "requires some other contributor' is changed to explicitly allow "
    "model misspecification and measurement error (Section 3.8). Code "
    "availability is upgraded from 'on request' to stating intent to "
    "publish a repository alongside submission. The Monte Carlo interval "
    "is retained but explicitly framed as secondary to the structural "
    "sensitivity analyses (Section 2.8, 3.7)."
)

d.save("Response_to_Reviewer_Round3.docx")
print("Saved: Response_to_Reviewer_Round3.docx")
